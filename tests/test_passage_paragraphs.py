"""Her paragraph breaks, on the page she prints.

`add_reading_passage` splits a passage on a blank line and only on a blank
line. Measured at `generators/components.py:835`: a passage whose paragraphs
are separated by single newlines is drawn as one unbroken wall of text — on
reading comprehension and problem solving, the two sheet types that print
prose whole and the two this feature rests on.

⚠️ **Normalising her pasted text is not enough.** `source_from_text` already
fixes what *she* supplies, and that covers reproduction. It does not cover the
one action where the printed passage is not her text at all: *simplify or
adapt* prints the model's own rewrite, which never went through it. The plan
listed this as a pre-existing defect; the adapt action is what makes it live.
"""

from docx import Document

from generators.components import add_reading_passage


def paragraphs_of(passage_text):
    """The paragraphs the child would actually see, out of a real document."""
    document = Document()
    add_reading_passage(document, {"title": "A passage", "text": passage_text}, "classic")
    table = document.tables[0]
    return [p.text for p in table.cell(0, 0).paragraphs if p.text.strip()]


class TestAPassageIsDrawnWithItsBreaks:
    def test_blank_lines_are_paragraph_breaks(self):
        assert len(paragraphs_of("First para.\n\nSecond para.")) == 2

    def test_single_newlines_are_paragraph_breaks_too(self):
        """The defect. A rewritten passage using single newlines was drawn as
        one block, however many paragraphs the model wrote."""
        assert len(paragraphs_of("First para.\nSecond para.\nThird para.")) == 3

    def test_a_run_of_blank_lines_is_still_one_break(self):
        assert len(paragraphs_of("First.\n\n\n\nSecond.")) == 2

    def test_a_passage_with_no_breaks_is_one_paragraph(self):
        assert len(paragraphs_of("One sentence and then another one.")) == 1

    def test_the_words_all_survive(self):
        """The control: a split that drops text would pass the counts above."""
        drawn = " ".join(paragraphs_of("First para.\nSecond para.\nThird para."))
        assert drawn == "First para. Second para. Third para."

    def test_wrapped_lines_inside_a_paragraph_are_not_broken_apart(self):
        """When she has used blank lines, single newlines are soft wrapping —
        breaking on those would shatter one paragraph into a dozen."""
        wrapped = "A sentence that was\nwrapped by the editor.\n\nA second paragraph."
        assert len(paragraphs_of(wrapped)) == 2
